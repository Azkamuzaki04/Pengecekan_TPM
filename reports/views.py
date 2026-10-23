from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render
from django.utils import timezone
from inspection.models import LaporanInspeksi
from accounts.views import get_role


def _filtered_qs(request):
    qs = LaporanInspeksi.objects.select_related("barang", "user", "approved_by").prefetch_related("hasil_parameter__parameter").order_by("-tanggal_inspeksi")
    if get_role(request.user) == "PETUGAS":
        qs = qs.filter(user=request.user)
    q = request.GET.get("q", "").strip()
    hasil = request.GET.get("hasil", "").strip()
    approval = request.GET.get("approval", "").strip()
    start = request.GET.get("start", "").strip()
    end = request.GET.get("end", "").strip()
    if q: qs = qs.filter(nomor_laporan__icontains=q) | qs.filter(barang__nama_barang__icontains=q) | qs.filter(barang__barcode__icontains=q) | qs.filter(user__username__icontains=q)
    if hasil in {"OK", "NG"}: qs = qs.filter(status_hasil=hasil)
    if approval in {"PENDING", "APPROVED", "REJECTED"}: qs = qs.filter(status_approval=approval)
    if start: qs = qs.filter(tanggal_inspeksi__date__gte=start)
    if end: qs = qs.filter(tanggal_inspeksi__date__lte=end)
    return qs.distinct()


@login_required
def report_list(request):
    qs = _filtered_qs(request)
    context = {
        "laporan_list": qs[:300], "total": qs.count(),
        "ok_count": qs.filter(status_hasil="OK").count(), "ng_count": qs.filter(status_hasil="NG").count(),
        "pending_count": qs.filter(status_approval="PENDING").count(),
    }
    return render(request, "reports/list.html", context)


@login_required
def export_excel(request):
    if get_role(request.user) != "ADMIN":
        return HttpResponseForbidden("Export Excel hanya untuk administrator.")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    qs = _filtered_qs(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan"
    headers = ["No. Laporan", "Tanggal", "Kode Barang", "Nama Barang", "Barcode", "Petugas", "Hasil", "Approval", "Disetujui Oleh", "Waktu Approval", "Catatan"]
    ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="C8102E"); c.font = Font(bold=True, color="FFFFFF")
    for item in qs:
        ws.append([item.nomor_laporan, item.tanggal_inspeksi.strftime("%Y-%m-%d %H:%M"), item.barang.kode_barang, item.barang.nama_barang, item.barang.barcode, item.user.username, item.status_hasil, item.get_status_approval_display(), item.approved_by.username if item.approved_by else "", item.approved_at.strftime("%Y-%m-%d %H:%M") if item.approved_at else "", item.catatan or ""])
    detail = wb.create_sheet("Detail Pemeriksaan")
    detail_headers = ["No. Laporan", "Tanggal", "Barang", "Petugas", "No", "Item Pemeriksaan", "Instruksi", "Standar", "Nilai/Hasil", "Status", "Catatan"]
    detail.append(detail_headers)
    for c in detail[1]: c.fill = PatternFill("solid", fgColor="333333"); c.font = Font(bold=True, color="FFFFFF")
    for item in qs:
        for h in item.hasil_parameter.all():
            detail.append([item.nomor_laporan, item.tanggal_inspeksi.strftime("%Y-%m-%d %H:%M"), item.barang.nama_barang, item.user.username, h.parameter.urutan, h.parameter.nama_parameter, h.parameter.instruksi or "", h.parameter.standar or "", h.nilai_hasil or "", h.get_hasil_display(), h.catatan or ""])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 45)
        for row in sheet.iter_rows():
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="checksit-laporan-{timezone.now():%Y%m%d-%H%M}.xlsx"'
    wb.save(response)
    return response


# Backward-compatible route name used by older templates.
export_csv = export_excel
